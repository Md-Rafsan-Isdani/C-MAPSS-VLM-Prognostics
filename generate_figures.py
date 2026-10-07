# =============================================================================
# generate_figures.py
#
# Regenerates all ten figures of the manuscript
#   "Reliability, Abstention and Operating-Regime Sensitivity of Vision-Language
#    Model Evidence in Multimodal Prognostics: A Multi-Seed Study on NASA C-MAPSS"
#
# >>> RUN THIS ON GOOGLE COLAB: paste the whole file into ONE Colab cell and run. <<<
#     No GPU is needed (CPU runtime is fine). Output: /content/paper_figures/
#     (300-dpi PNG + vector PDF for every figure), zipped and downloaded at the end.
#
# DATA
#   By default the figures are drawn from PAPER_DATA below: the values logged by the
#   three runs of vlm_prognostics.ipynb (seeds 42, 1729, 191; NVIDIA L4, PyTorch 2.11).
#   Every plotted value also appears in the manuscript's tables.
#   To rebuild PAPER_DATA from the run logs instead, set REBUILD_FROM_LOGS = True and
#   upload the three results files exported from the notebook output when prompted
#   (Word files containing the full printed log of each run; seed is read from the name).
#
# Figure -> manuscript mapping (file names match \registerfig in the .tex):
#   fig0_system_flowchart.png        Fig. 1   system flowchart (two-branch framework)
#   fig1_rmse_all_methods.png        Fig. 2   test RMSE, all methods, all seeds
#   fig9_seed_skill_radar.png        Fig. 3   relative skill radar per subset
#   fig2_vlm_contribution.png        Fig. 4   RMSE change from VLM evidence
#   fig5_help_hurt.png               Fig. 5   per-engine helped / hurt / no effect
#   fig3_evidence_quality.png        Fig. 6   VLM evidence quality
#   fig4_regime_association.png      Fig. 7   unreadability vs operating regimes
#   fig8_evidence_radar.png          Fig. 8   evidence profile radar
#   fig6_agreement.png               Fig. 9   VLM-numerical agreement (seed 42)
#   fig7_confidence_reliability.png  Fig. 10  confidence score vs accuracy (seed 42)
# =============================================================================

REBUILD_FROM_LOGS = False          # True -> parse the three uploaded results .docx logs
OUT = "/content/paper_figures"     # output folder (Colab)

import json, math, os, re, ast, zipfile
import numpy as np
import matplotlib
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from matplotlib.patches import Patch

# ----------------------------------------------------------------------------- data
PAPER_DATA = json.loads(r"""{"42":{"FD001":{"rmse":{"Linear regression":16.001,"Random forest":12.495,"XGBoost":12.274,"BiLSTM":13.045,"TCN":14.691,"MS-TCN":14.468,"MS-TCN + BiLSTM":13.656,"Transformer":14.527,"Simple average":12.601,"Weighted average":12.59,"Fusion, numerical only":12.935,"Fusion with VLM evidence":13.016},"n":100,"helped":32,"hurt":47,"no_effect":21,"mean_change":0.0426,"agree":56,"conflict":24,"vlm_abstained":20,"quality":{"clear":69,"unreadable":20,"weak":11},"terminal_pct":26.0,"unreadable_pct":20.0,"conf_bins":[[0.311,1.0,3],[0.454,1.0,4],[0.679,0.8,30],[0.936,0.73,63]],"conf_gap":0.209,"vlm_contribution":-0.081},"FD002":{"rmse":{"Linear regression":18.401,"Random forest":14.215,"XGBoost":13.483,"BiLSTM":13.601,"TCN":15.681,"MS-TCN":15.613,"MS-TCN + BiLSTM":14.229,"Transformer":15.723,"Simple average":13.801,"Weighted average":13.769,"Fusion, numerical only":13.66,"Fusion with VLM evidence":13.65},"n":259,"helped":53,"hurt":52,"no_effect":154,"mean_change":-0.0082,"agree":41,"conflict":65,"vlm_abstained":153,"quality":{"unreadable":154,"clear":92,"weak":13},"terminal_pct":10.0,"unreadable_pct":59.5,"conf_bins":[[0.19,1.0,1],[0.325,1.0,10],[0.542,1.0,32],[0.663,0.703,158],[0.915,0.655,58]],"conf_gap":0.168,"vlm_contribution":0.011},"FD003":{"rmse":{"Linear regression":16.906,"Random forest":12.375,"XGBoost":12.328,"BiLSTM":13.474,"TCN":18.959,"MS-TCN":12.998,"MS-TCN + BiLSTM":13.026,"Transformer":13.387,"Simple average":12.271,"Weighted average":12.091,"Fusion, numerical only":12.024,"Fusion with VLM evidence":12.146},"n":100,"helped":38,"hurt":46,"no_effect":16,"mean_change":0.0624,"agree":63,"conflict":24,"vlm_abstained":13,"quality":{"clear":81,"unreadable":13,"weak":6},"terminal_pct":33.0,"unreadable_pct":13.0,"conf_bins":[[0.313,1.0,2],[0.51,1.0,3],[0.689,0.833,24],[0.924,0.746,71]],"conf_gap":0.189,"vlm_contribution":-0.122},"FD004":{"rmse":{"Linear regression":21.317,"Random forest":18.274,"XGBoost":16.111,"BiLSTM":16.007,"TCN":16.798,"MS-TCN":15.285,"MS-TCN + BiLSTM":15.836,"Transformer":15.655,"Simple average":15.403,"Weighted average":15.328,"Fusion, numerical only":15.287,"Fusion with VLM evidence":15.31},"n":248,"helped":33,"hurt":71,"no_effect":144,"mean_change":0.0321,"agree":47,"conflict":58,"vlm_abstained":143,"quality":{"unreadable":144,"clear":91,"weak":13},"terminal_pct":9.7,"unreadable_pct":58.1,"conf_bins":[[0.307,1.0,1],[0.532,0.963,27],[0.666,0.748,163],[0.923,0.702,57]],"conf_gap":0.155,"vlm_contribution":-0.023}},"1729":{"FD001":{"rmse":{"Linear regression":15.891,"Random forest":12.42,"XGBoost":12.054,"BiLSTM":12.619,"TCN":14.597,"MS-TCN":13.527,"MS-TCN + BiLSTM":12.971,"Transformer":13.071,"Simple average":12.146,"Weighted average":12.104,"Fusion, numerical only":12.133,"Fusion with VLM evidence":12.214},"n":100,"helped":31,"hurt":48,"no_effect":21,"mean_change":0.0568,"agree":57,"conflict":23,"vlm_abstained":20,"quality":{"clear":69,"unreadable":20,"weak":11},"terminal_pct":26.0,"unreadable_pct":20.0,"conf_bins":[[0.331,1.0,3],[0.478,1.0,4],[0.684,0.833,30],[0.935,0.778,63]],"conf_gap":0.185,"vlm_contribution":-0.081},"FD002":{"rmse":{"Linear regression":18.278,"Random forest":13.922,"XGBoost":13.469,"BiLSTM":13.444,"TCN":14.857,"MS-TCN":14.107,"MS-TCN + BiLSTM":13.858,"Transformer":14.04,"Simple average":13.2,"Weighted average":13.153,"Fusion, numerical only":13.052,"Fusion with VLM evidence":13.061},"n":259,"helped":53,"hurt":53,"no_effect":153,"mean_change":0.0058,"agree":40,"conflict":66,"vlm_abstained":153,"quality":{"unreadable":154,"clear":92,"weak":13},"terminal_pct":10.0,"unreadable_pct":59.5,"conf_bins":[[0.19,1.0,1],[0.318,1.0,7],[0.533,0.971,35],[0.665,0.75,156],[0.916,0.717,60]],"conf_gap":0.178,"vlm_contribution":-0.009},"FD003":{"rmse":{"Linear regression":16.422,"Random forest":13.516,"XGBoost":13.043,"BiLSTM":13.915,"TCN":16.028,"MS-TCN":12.536,"MS-TCN + BiLSTM":12.914,"Transformer":14.191,"Simple average":12.366,"Weighted average":12.286,"Fusion, numerical only":12.45,"Fusion with VLM evidence":12.561},"n":100,"helped":33,"hurt":53,"no_effect":14,"mean_change":0.0767,"agree":62,"conflict":25,"vlm_abstained":13,"quality":{"clear":81,"unreadable":13,"weak":6},"terminal_pct":33.0,"unreadable_pct":13.0,"conf_bins":[[0.314,1.0,1],[0.482,1.0,2],[0.675,0.826,23],[0.925,0.73,74]],"conf_gap":0.197,"vlm_contribution":-0.111},"FD004":{"rmse":{"Linear regression":21.166,"Random forest":16.635,"XGBoost":15.921,"BiLSTM":15.478,"TCN":16.737,"MS-TCN":15.751,"MS-TCN + BiLSTM":16.212,"Transformer":15.239,"Simple average":15.219,"Weighted average":15.142,"Fusion, numerical only":15.256,"Fusion with VLM evidence":15.265},"n":248,"helped":35,"hurt":69,"no_effect":144,"mean_change":0.0322,"agree":48,"conflict":57,"vlm_abstained":143,"quality":{"unreadable":144,"clear":91,"weak":13},"terminal_pct":9.7,"unreadable_pct":58.1,"conf_bins":[[0.312,1.0,2],[0.526,0.939,33],[0.665,0.719,153],[0.918,0.7,60]],"conf_gap":0.147,"vlm_contribution":-0.009}},"191":{"FD001":{"rmse":{"Linear regression":16.058,"Random forest":12.431,"XGBoost":12.104,"BiLSTM":13.32,"TCN":15.056,"MS-TCN":14.081,"MS-TCN + BiLSTM":13.305,"Transformer":13.325,"Simple average":12.368,"Weighted average":12.337,"Fusion, numerical only":12.375,"Fusion with VLM evidence":12.402},"n":100,"helped":35,"hurt":45,"no_effect":20,"mean_change":-0.002,"agree":55,"conflict":25,"vlm_abstained":20,"quality":{"clear":69,"unreadable":20,"weak":11},"terminal_pct":26.0,"unreadable_pct":20.0,"conf_bins":[[0.318,1.0,3],[0.462,1.0,4],[0.679,0.833,30],[0.933,0.746,63]],"conf_gap":0.206,"vlm_contribution":-0.027},"FD002":{"rmse":{"Linear regression":18.346,"Random forest":14.1,"XGBoost":13.256,"BiLSTM":13.834,"TCN":14.727,"MS-TCN":14.962,"MS-TCN + BiLSTM":14.209,"Transformer":14.652,"Simple average":13.444,"Weighted average":13.418,"Fusion, numerical only":13.411,"Fusion with VLM evidence":13.407},"n":259,"helped":55,"hurt":50,"no_effect":154,"mean_change":-0.005,"agree":41,"conflict":65,"vlm_abstained":153,"quality":{"unreadable":154,"clear":92,"weak":13},"terminal_pct":10.0,"unreadable_pct":59.5,"conf_bins":[[0.19,1.0,4],[0.331,1.0,7],[0.542,1.0,34],[0.664,0.724,156],[0.916,0.655,58]],"conf_gap":0.185,"vlm_contribution":0.004},"FD003":{"rmse":{"Linear regression":16.431,"Random forest":13.302,"XGBoost":12.545,"BiLSTM":13.388,"TCN":19.119,"MS-TCN":12.5,"MS-TCN + BiLSTM":12.769,"Transformer":13.16,"Simple average":12.325,"Weighted average":12.179,"Fusion, numerical only":12.047,"Fusion with VLM evidence":12.232},"n":100,"helped":42,"hurt":44,"no_effect":14,"mean_change":0.0882,"agree":60,"conflict":27,"vlm_abstained":13,"quality":{"clear":81,"unreadable":13,"weak":6},"terminal_pct":33.0,"unreadable_pct":13.0,"conf_bins":[[0.307,1.0,2],[0.486,1.0,2],[0.693,0.923,26],[0.924,0.743,70]],"conf_gap":0.211,"vlm_contribution":-0.185},"FD004":{"rmse":{"Linear regression":21.124,"Random forest":17.247,"XGBoost":16.404,"BiLSTM":17.08,"TCN":19.755,"MS-TCN":15.713,"MS-TCN + BiLSTM":15.398,"Transformer":15.058,"Simple average":15.503,"Weighted average":15.366,"Fusion, numerical only":15.676,"Fusion with VLM evidence":15.715},"n":248,"helped":35,"hurt":70,"no_effect":143,"mean_change":0.0555,"agree":48,"conflict":57,"vlm_abstained":143,"quality":{"unreadable":144,"clear":91,"weak":13},"terminal_pct":9.7,"unreadable_pct":58.1,"conf_bins":[[0.327,1.0,4],[0.543,0.931,29],[0.665,0.694,157],[0.923,0.672,58]],"conf_gap":0.133,"vlm_contribution":-0.039}}}""")


def parse_results_docx(path):
    """Extract the per-subset values used by the figures from one results .docx log
    (the notebook's printed output saved as a Word file)."""
    import docx                                          # pip install python-docx
    d = docx.Document(path); T = d.tables
    paras = "\n".join(p.text for p in d.paragraphs)
    names = {"LinearRegression": "Linear regression", "RandomForest": "Random forest", "XGBoost": "XGBoost",
             "BiLSTM": "BiLSTM", "TCN": "TCN", "MS-TCN": "MS-TCN", "MS-TCN+BiLSTM": "MS-TCN + BiLSTM",
             "Transformer": "Transformer", "SimpleAvg_Ensemble": "Simple average",
             "WeightedAvg_Ensemble": "Weighted average", "ConstrainedFusion (numeric only)": "Fusion, numerical only",
             "ConstrainedFusion+VLM (proposed)": "Fusion with VLM evidence"}
    out = {}
    for k, s in enumerate(["FD001", "FD002", "FD003", "FD004"]):
        rmse = {}
        for row in T[2 + 5 * k].rows[1:]:                # metrics table of subset k
            c = [x.text.strip() for x in row.cells]
            if c[0] in names: rmse[names[c[0]]] = float(c[1])
        conf = []
        for row in T[4 + 5 * k].rows[1:]:                # confidence-reliability table
            c = [x.text.strip() for x in row.cells]
            if c[1] not in ("NaN", ""): conf.append([float(c[1]), float(c[2]), int(c[3])])
        i = paras.index(f"[TITLE] {s} — VLM evidence diagnostics"); blk = paras[i:]
        g = lambda pat: re.search(pat, blk, re.M).group(1)
        pre = paras[:i]
        e = dict(rmse=rmse, conf_bins=conf,
                 helped=int(g(r"^\s*helped\s+(\d+)")), hurt=int(g(r"^\s*hurt\s+(\d+)")),
                 no_effect=int(g(r"^\s*no_effect\s+(\d+)")),
                 mean_change=float(g(r"when VLM evidence is added: ([+-][\d.]+)")),
                 agree=int(g(r"^\s*agree\s+(\d+)")), conflict=int(g(r"^\s*conflict\s+(\d+)")),
                 vlm_abstained=int(g(r"^\s*vlm_abstained\s+(\d+)")),
                 quality=ast.literal_eval(g(r"vlm_evidence_quality: (\{.*?\})")),
                 terminal_pct=float(g(r"reported for ([\d.]+)% of engines")),
                 unreadable_pct=float(g(r"judged unreadable: ([\d.]+)%")),
                 conf_gap=float(re.findall(r"reliability gap, \|err\|<=15 cycles\): ([\d.]+)", pre)[-1]))
        e["n"] = e["helped"] + e["hurt"] + e["no_effect"]
        out[s] = e
    for row in T[21].rows[1:]:                           # cross-subset summary: exact VLM contribution
        c = [x.text.strip() for x in row.cells]
        out[c[1]]["vlm_contribution"] = float(c[7])
    return out


if REBUILD_FROM_LOGS:
    os.system("pip -q install python-docx")
    from google.colab import files
    print("Upload the three results .docx files (seed 42, 1729 and 191):")
    up = files.upload()
    D = {}
    for fname in up:
        seed = next(int(x) for x in re.findall(r"\d+", fname) if int(x) in (42, 1729, 191))
        D[seed] = parse_results_docx(fname)
    assert sorted(D) == [42, 191, 1729], f"need all three seeds, got {sorted(D)}"
else:
    D = {int(k): v for k, v in PAPER_DATA.items()}

os.makedirs(OUT, exist_ok=True)

# ----------------------------------------------------------------------------- style
TEXTWIDTH = 6.0          # inches; Springer sn-jnl text block is ~6 in
SEEDS = [42, 1729, 191]
SUBSETS = ['FD001', 'FD002', 'FD003', 'FD004']
REGIME_ORDER = ['FD001', 'FD003', 'FD002', 'FD004']     # one regime | six regimes

# ---- palette (validated: dataviz reference palette, light mode, white surface) ----
INK, INK2, MUTED = '#0b0b0b', '#52514e', '#898781'
GRID, AXIS = '#e6e5df', '#8f8e88'
SEED_C = {42: '#2a78d6', 1729: '#eb6834', 191: '#1baf7a'}           # categorical slots 1-3
SEED_M = {42: 'o', 1729: 's', 191: '^'}                                 # secondary encoding
SUB_C = {'FD001': '#2a78d6', 'FD003': '#1baf7a', 'FD002': '#eb6834', 'FD004': '#4a3aa7'}
SUB_M = {'FD001': 'o', 'FD003': 's', 'FD002': '^', 'FD004': 'D'}
SUB_LS = {'FD001': '-', 'FD003': '-', 'FD002': (0, (4, 2)), 'FD004': (0, (4, 2))}
GOOD, WARN, SERIOUS, CRIT, NEUTRAL = '#0ca30c', '#fab219', '#ec835a', '#d03b3b', '#c9c8c2'
REGIME_C = {1: '#2a78d6', 6: '#eb6834'}

plt.rcParams.update({
    'font.family': 'DejaVu Sans', 'font.size': 8, 'axes.titlesize': 9, 'axes.labelsize': 8,
    'xtick.labelsize': 7.5, 'ytick.labelsize': 7.5, 'legend.fontsize': 7.5,
    'text.color': INK, 'axes.labelcolor': INK, 'xtick.color': INK2, 'ytick.color': INK2,
    'axes.edgecolor': AXIS, 'axes.linewidth': 0.6, 'xtick.major.width': 0.6, 'ytick.major.width': 0.6,
    'xtick.major.size': 2.5, 'ytick.major.size': 2.5, 'axes.spines.top': False, 'axes.spines.right': False,
    'axes.grid': False, 'grid.color': GRID, 'grid.linewidth': 0.5, 'legend.frameon': False,
    'figure.facecolor': 'white', 'axes.facecolor': 'white', 'savefig.facecolor': 'white',
    'axes.titleweight': 'bold', 'axes.titlepad': 6,
})


def save(fig, name):
    fig.savefig(os.path.join(OUT, name), dpi=300, bbox_inches="tight", pad_inches=0.03)
    fig.savefig(os.path.join(OUT, name.replace(".png", ".pdf")), bbox_inches="tight", pad_inches=0.03)
    plt.close(fig)
    print("wrote", name)


# ----------------------------------------------------------------------------- figures
def pct(n, d): return 100.0 * n / d

def regime_axis(ax, labels_y=-0.2):
    """x ticks in REGIME_ORDER with a divider and group labels under the axis."""
    ax.set_xticks(range(4)); ax.set_xticklabels(REGIME_ORDER, fontweight='bold', color=INK)
    ax.axvline(1.5, color=MUTED, lw=0.7, ls=(0, (3, 2)), zorder=0)
    for x, t in [(0.5, 'One regime'), (2.5, 'Six regimes')]:
        ax.text(x, labels_y, t, transform=ax.get_xaxis_transform(), ha='center', va='top',
                fontsize=7, color=INK2)
    ax.tick_params(axis='x', length=0)

def stacked(ax, cats, colours, text_colours, values, width=0.66, fs=6.8):
    """values: list over REGIME_ORDER of lists (one per category) in percent."""
    for i, vals in enumerate(values):
        bottom = 0
        for v, c, tc in zip(vals, colours, text_colours):
            ax.bar(i, v, width, bottom=bottom, color=c, edgecolor='white', linewidth=0.8, zorder=2)
            if v >= 4.5:
                ax.text(i, bottom + v / 2, f'{v:.1f}', ha='center', va='center', fontsize=fs,
                        color=tc, fontweight='bold', zorder=3)
            bottom += v
    ax.set_ylim(0, 100); ax.set_yticks(range(0, 101, 20)); ax.set_xlim(-0.55, 3.55)
    ax.yaxis.grid(True); ax.set_axisbelow(True)


# ------------------------------------------------------------------ Fig 1
def fig_rmse():
    methods = ['Linear regression', 'Random forest', 'XGBoost', 'BiLSTM', 'TCN', 'MS-TCN',
               'MS-TCN + BiLSTM', 'Transformer', 'Simple average', 'Weighted average',
               'Fusion, numerical only', 'Fusion with VLM evidence']
    ypos = {m: i + (0.6 if i >= 8 else 0) for i, m in enumerate(methods)}
    fig, axes = plt.subplots(2, 2, figsize=(0.9 * TEXTWIDTH, 6.0), sharey=True)
    off = {42: -0.24, 1729: 0.0, 191: 0.24}
    for ax, s in zip(axes.flat, SUBSETS):
        allv = [D[sd][s]['rmse'][m] for sd in SEEDS for m in methods]
        lo, hi = math.floor(min(allv) - 0.3), math.ceil(max(allv) + 0.3)
        for m in methods:
            v = [D[sd][s]['rmse'][m] for sd in SEEDS]
            ax.plot([min(v), max(v)], [ypos[m]] * 2, color='#d4d3cc', lw=3.2, solid_capstyle='round', zorder=1)
            for sd in SEEDS:
                ax.plot(D[sd][s]['rmse'][m], ypos[m] + off[sd], marker=SEED_M[sd], ms=4.2, ls='none',
                        color=SEED_C[sd], mec='white', mew=0.5, zorder=3)
        ax.axhline(7.8, color=AXIS, lw=0.6, ls=(0, (2, 2)))
        ax.set_xlim(lo, hi); ax.set_xticks(range(lo, hi + 1))
        ax.xaxis.grid(True); ax.set_axisbelow(True)
        ax.set_title(s, loc='left')
        ax.tick_params(axis='y', length=0)
    for ax in axes[1]: ax.set_xlabel('Test RMSE (cycles)')
    axes[0, 0].set_yticks([ypos[m] for m in methods])
    for ax in axes[:, 0]: ax.set_yticklabels(methods, color=INK)
    axes[0, 0].set_ylim(12.2, -0.6)
    handles = [Line2D([], [], marker=SEED_M[sd], color=SEED_C[sd], ls='none', ms=5, mec='white', mew=0.5,
                      label=f'Seed {sd}') for sd in SEEDS]
    handles.append(Line2D([], [], color='#d4d3cc', lw=3.2, label='Range across seeds'))
    fig.legend(handles=handles, loc='lower center', ncol=4, bbox_to_anchor=(0.55, -0.025))
    fig.tight_layout(rect=(0, 0.03, 1, 1), h_pad=1.6, w_pad=1.2)
    save(fig, 'fig1_rmse_all_methods.png')


# ------------------------------------------------------------------ Fig 2 (skill radar)
def fig_skill():
    methods = ['XGBoost', 'BiLSTM', 'MS-TCN', 'Transformer', 'Weighted average', 'Fusion with VLM evidence']
    labels = ['XGBoost', 'BiLSTM', 'MS-TCN', 'Transformer', 'Weighted\naverage', 'Fusion with\nVLM evidence']
    allm = list(D[42]['FD001']['rmse'])
    ang = np.linspace(0, 2 * np.pi, len(methods), endpoint=False)
    fig, axes = plt.subplots(2, 2, figsize=(0.9 * TEXTWIDTH, 5.0), subplot_kw={'polar': True})
    for ax, s in zip(axes.flat, SUBSETS):
        pool = [D[sd][s]['rmse'][m] for sd in SEEDS for m in allm]
        best, worst = min(pool), max(pool)
        for sd in SEEDS:
            v = [(worst - D[sd][s]['rmse'][m]) / (worst - best) for m in methods]
            ax.plot(np.r_[ang, ang[0]], v + v[:1], color=SEED_C[sd], lw=1.5, zorder=3)
            ax.plot(ang, v, ls='none', marker=SEED_M[sd], ms=3.6, color=SEED_C[sd], mec='white', mew=0.4, zorder=4)
        ax.set_theta_offset(np.pi / 2); ax.set_theta_direction(-1)
        ax.set_ylim(0, 1.0); ax.set_yticks([0.25, 0.5, 0.75, 1.0]); ax.set_yticklabels([])
        ax.set_xticks(ang); ax.set_xticklabels(labels, fontsize=7, color=INK)
        ax.tick_params(axis='x', pad=9)
        ax.grid(color=GRID, lw=0.6); ax.spines['polar'].set_color(AXIS); ax.spines['polar'].set_linewidth(0.8)
        ax.set_title(s, pad=14)
    handles = [Line2D([], [], color=SEED_C[sd], lw=1.5, marker=SEED_M[sd], ms=4, mec='white', mew=0.4,
                      label=f'Seed {sd}') for sd in SEEDS]
    fig.legend(handles=handles, loc='lower center', ncol=3, bbox_to_anchor=(0.5, -0.01))
    fig.tight_layout(rect=(0, 0.04, 1, 1), h_pad=2.6, w_pad=4.0)
    save(fig, 'fig9_seed_skill_radar.png')


# ------------------------------------------------------------------ Fig 3 (VLM contribution)
def fig_contrib():
    fig, ax = plt.subplots(figsize=(0.9 * TEXTWIDTH, 3.0))
    w = 0.26
    for j, sd in enumerate(SEEDS):
        for i, s in enumerate(SUBSETS):
            v = D[sd][s]['vlm_contribution']          # numerical-only RMSE minus RMSE with VLM evidence
            x = i + (j - 1) * w
            ax.bar(x, v, w * 0.9, color=SEED_C[sd], zorder=2, label=f'Seed {sd}' if i == 0 else None)
            ax.text(x, v + (0.005 if v >= 0 else -0.005), f'{v:+.3f}'.replace('-', '\u2212'), rotation=90,
                    ha='center', va='bottom' if v >= 0 else 'top', fontsize=6.8, color=INK)
    ax.axhline(0, color=INK, lw=0.8, zorder=3)
    ax.set_xticks(range(4)); ax.set_xticklabels(SUBSETS, fontweight='bold', color=INK); ax.tick_params(axis='x', length=0)
    ax.set_xlim(-0.55, 3.55); ax.set_ylim(-0.25, 0.05)
    ticks = [-0.20, -0.15, -0.10, -0.05, 0.0]
    ax.set_yticks(ticks); ax.set_yticklabels([f'{t:.2f}'.replace('-', '\u2212') if t else '0' for t in ticks])
    ax.yaxis.grid(True); ax.set_axisbelow(True)
    ax.set_ylabel('RMSE reduction (cycles)\n(below zero: VLM evidence increased error)', linespacing=1.4)
    ax.legend(loc='lower center', ncol=3, bbox_to_anchor=(0.5, 1.0))
    fig.tight_layout()
    save(fig, 'fig2_vlm_contribution.png')


# ------------------------------------------------------------------ Fig 4 (help / hurt)
def fig_helphurt():
    fig, axes = plt.subplots(1, 3, figsize=(0.9 * TEXTWIDTH, 2.9), sharey=True)
    for ax, sd in zip(axes, SEEDS):
        vals = []
        for s in REGIME_ORDER:
            d = D[sd][s]; n = d['n']
            vals.append([pct(d['helped'], n), pct(d['hurt'], n), pct(d['no_effect'], n)])
        stacked(ax, None, [GOOD, CRIT, NEUTRAL], ['white', 'white', INK], vals, width=0.8, fs=6.4)
        regime_axis(ax, labels_y=-0.12)
        for t in ax.texts:
            if t.get_text().startswith('One'): t.set_fontsize(6.6)
            if t.get_text().startswith('Six'): t.set_fontsize(6.6)
        ax.set_xticklabels(REGIME_ORDER, fontsize=6.4, fontweight='bold', color=INK)
        ax.set_xlim(-0.48, 3.48)
        ax.set_title(f'Seed {sd}')
    axes[0].set_ylabel('Test engines (%)')
    for ax in axes[1:]: ax.tick_params(axis='y', length=0)
    handles = [Patch(color=c, label=l) for c, l in [(GOOD, 'Helped (error fell)'), (CRIT, 'Hurt (error rose)'),
                                                   (NEUTRAL, 'No effect')]]
    fig.legend(handles=handles, loc='lower center', ncol=3, bbox_to_anchor=(0.53, -0.02))
    fig.tight_layout(rect=(0, 0.07, 1, 1), w_pad=0.6)
    save(fig, 'fig5_help_hurt.png')


# ------------------------------------------------------------------ Fig 5 (evidence quality)
def fig_quality():
    fig, ax = plt.subplots(figsize=(0.6 * TEXTWIDTH, 2.9))
    vals = []
    for s in REGIME_ORDER:
        q = D[42][s]['quality']; n = sum(q.values())
        vals.append([pct(q['clear'], n), pct(q['weak'], n), pct(q['unreadable'], n)])
    stacked(ax, None, [GOOD, WARN, CRIT], ['white', INK, 'white'], vals)
    regime_axis(ax, labels_y=-0.1)
    ax.set_ylabel('Test engines (%)')
    handles = [Patch(color=c, label=l) for c, l in [(GOOD, 'Clear'), (WARN, 'Weak'), (CRIT, 'Unreadable')]]
    ax.legend(handles=handles, loc='upper center', bbox_to_anchor=(0.5, -0.2), ncol=3)
    fig.tight_layout()
    save(fig, 'fig3_evidence_quality.png')


# ------------------------------------------------------------------ Fig 6 (regime association)
def fig_regime():
    fig, (a, b) = plt.subplots(1, 2, figsize=(0.9 * TEXTWIDTH, 2.7), gridspec_kw={'width_ratios': [1, 1.25]})
    regimes = {'FD001': 1, 'FD003': 1, 'FD002': 6, 'FD004': 6}
    pos = {'FD001': (-0.07, 'right'), 'FD003': (0.07, 'left'), 'FD002': (0.93, 'right'), 'FD004': (1.07, 'left')}
    for s in REGIME_ORDER:
        x, ha = pos[s]; y = D[42][s]['unreadable_pct']
        a.plot(x, y, 'o', ms=7, color=REGIME_C[regimes[s]], mec='white', mew=0.8, zorder=3)
        a.text(x + (-0.06 if ha == 'right' else 0.06), y, f'{s}\n{y:.1f}', ha=ha, va='center', fontsize=7,
               color=INK, linespacing=1.2)
    a.set_xlim(-0.6, 1.6); a.set_xticks([0, 1]); a.set_xticklabels(['1', '6'], fontweight='bold', color=INK)
    a.set_xlabel('Number of operating regimes')
    a.set_ylim(0, 75); a.set_yticks(range(0, 71, 10)); a.yaxis.grid(True); a.set_axisbelow(True)
    a.set_ylabel('Figures judged unreadable (%)'); a.tick_params(axis='x', length=0)
    a.set_title('(a)', loc='left')

    w = 0.36
    for i, s in enumerate(REGIME_ORDER):
        for k, (val, col) in enumerate([(D[42][s]['unreadable_pct'], CRIT), (D[42][s]['terminal_pct'], SEED_C[42])]):
            x = i + (k - 0.5) * w
            b.bar(x, val, w * 0.92, color=col, zorder=2)
            b.text(x, val + 1.2, f'{val:.1f}', ha='center', va='bottom', fontsize=6.6, color=INK)
    regime_axis(b, labels_y=-0.12)
    b.set_ylim(0, 75); b.set_yticks(range(0, 71, 10)); b.yaxis.grid(True); b.set_axisbelow(True)
    b.set_ylabel('Test engines (%)'); b.set_title('(b)', loc='left')
    for t in b.texts:
        if t.get_text().startswith('One'): t.set_text('One regime')
        if t.get_text().startswith('Six'): t.set_text('Six regimes')
    fig.legend(handles=[Patch(color=CRIT, label='Unreadable'), Patch(color=SEED_C[42], label='Terminal phase reported')],
               loc='lower center', bbox_to_anchor=(0.5, -0.01), ncol=2)
    fig.tight_layout(rect=(0, 0.07, 1, 1), w_pad=2.0)
    save(fig, 'fig4_regime_association.png')


# ------------------------------------------------------------------ Fig 7 (evidence radar)
def fig_radar():
    axes_lbl = ['Terminal phase\nreported', 'Clear\nevidence', 'Readable\nfigure', 'Agreement with\nnumerical\nmodels']
    ang = np.linspace(0, 2 * np.pi, 4, endpoint=False)
    fig = plt.figure(figsize=(0.48 * TEXTWIDTH, 3.0))
    ax = fig.add_axes([0.24, 0.26, 0.52, 0.52], polar=True)
    for s in REGIME_ORDER:
        d = D[42][s]; n = d['n']; q = d['quality']
        v = [d['terminal_pct'], pct(q['clear'], n), pct(q['clear'] + q['weak'], n), pct(d['agree'], n)]
        ax.plot(np.r_[ang, ang[0]], v + v[:1], color=SUB_C[s], ls=SUB_LS[s], lw=1.4, zorder=3)
        ax.plot(ang, v, ls='none', marker=SUB_M[s], ms=4, color=SUB_C[s], mec='white', mew=0.5, zorder=4)
    ax.set_theta_offset(np.pi / 2); ax.set_theta_direction(-1)
    ax.set_ylim(0, 100); ax.set_yticks([20, 40, 60, 80, 100])
    ax.set_yticklabels([])
    for r in [40, 60, 80, 100]:   # on the top axis, above every plotted value there (max 33%)
        ax.text(np.deg2rad(5), r, f'{r}%' if r == 100 else str(r), fontsize=5.8, color=INK2, ha='left',
                va='center', zorder=5, bbox=dict(boxstyle='round,pad=0.1', fc='white', ec='none', alpha=0.9))
    ax.set_xticks(ang); ax.set_xticklabels(axes_lbl, fontsize=6.8, color=INK)
    ax.tick_params(axis='x', pad=10)
    ax.grid(color=GRID, lw=0.6); ax.spines['polar'].set_color(AXIS)
    lbl = {'FD001': 'FD001 (one regime)', 'FD003': 'FD003 (one regime)',
           'FD002': 'FD002 (six regimes)', 'FD004': 'FD004 (six regimes)'}
    handles = [Line2D([], [], color=SUB_C[s], ls=SUB_LS[s], lw=1.4, marker=SUB_M[s], ms=4, mec='white', mew=0.5,
                      label=lbl[s]) for s in REGIME_ORDER]
    fig.legend(handles=handles, loc='lower center', ncol=2, fontsize=6.6, handlelength=2.4,
               bbox_to_anchor=(0.5, 0.02), columnspacing=1.0)
    save(fig, 'fig8_evidence_radar.png')


# ------------------------------------------------------------------ Fig 8 (agreement, seed 42)
def fig_agreement():
    fig, ax = plt.subplots(figsize=(0.48 * TEXTWIDTH, 2.9))
    vals = []
    for s in REGIME_ORDER:
        d = D[42][s]; n = d['n']
        vals.append([pct(d['agree'], n), pct(d['conflict'], n), pct(d['vlm_abstained'], n)])
    stacked(ax, None, [GOOD, SERIOUS, NEUTRAL], ['white', INK, INK], vals, width=0.7)
    regime_axis(ax, labels_y=-0.1)
    for t in ax.texts:
        if t.get_text().startswith('One'): t.set_text('One regime')
        if t.get_text().startswith('Six'): t.set_text('Six regimes')
    ax.set_xticklabels(REGIME_ORDER, fontsize=7, fontweight='bold', color=INK)
    ax.set_ylabel('Test engines (%)')
    handles = [Patch(color=c, label=l) for c, l in [(GOOD, 'Agree'), (SERIOUS, 'Conflict'), (NEUTRAL, 'VLM abstained')]]
    ax.legend(handles=handles, loc='upper center', bbox_to_anchor=(0.45, -0.2), ncol=3, fontsize=6.8,
              handlelength=1.2, columnspacing=0.9)
    fig.tight_layout()
    save(fig, 'fig6_agreement.png')


# ------------------------------------------------------------------ Fig 9 (confidence, seed 42)
def fig_confidence():
    fig, axes = plt.subplots(2, 2, figsize=(0.75 * TEXTWIDTH, 4.3), sharex=True, sharey=True)
    for ax, s in zip(axes.flat, SUBSETS):
        bins = [dict(mean_conf=a, acc=b, n=c) for a, b, c in D[42][s]['conf_bins']]
        x = [b['mean_conf'] for b in bins]; y = [b['acc'] for b in bins]; n = [b['n'] for b in bins]
        ax.plot([0, 1], [0, 1], color=MUTED, lw=0.9, ls=(0, (4, 3)), zorder=1)
        ax.plot(x, y, color=SEED_C[42], lw=1.5, zorder=2)
        ax.scatter(x, y, s=[12 + 2.2 * k for k in n], color=SEED_C[42], edgecolor='white', lw=0.7, zorder=3)
        ax.set_xlim(0, 1.0); ax.set_ylim(0.5, 1.06)
        ax.set_xticks([0, 0.2, 0.4, 0.6, 0.8, 1.0]); ax.set_yticks([0.5, 0.6, 0.7, 0.8, 0.9, 1.0])
        ax.grid(True); ax.set_axisbelow(True)
        ax.set_title(f'{s}', loc='left')
        if s in ('FD003', 'FD004'): ax.set_xlabel('Confidence score (mean of bin)')
        ax.text(0.98, 0.04, f'gap {D[42][s]["conf_gap"]:.3f}', transform=ax.transAxes, ha='right', va='bottom',
                fontsize=6.6, color=INK2)
    fig.supylabel('Observed accuracy (|error| ≤ 15 cycles)', fontsize=8, x=0.015)
    handles = [Line2D([], [], color=SEED_C[42], lw=1.5, marker='o', ms=4, mec='white', label='Observed (seed 42)'),
               Line2D([], [], color=MUTED, lw=0.9, ls=(0, (4, 3)), label='Ideal (score = accuracy)')]
    leg1 = fig.legend(handles=handles, loc='lower center', ncol=2, bbox_to_anchor=(0.53, 0.045))
    sizes = [Line2D([], [], ls='none', marker='o', color=SEED_C[42], mec='white', ms=math.sqrt(12 + 2.2 * k),
                    label=f'{k} engines') for k in (10, 50, 150)]
    fig.legend(handles=sizes, loc='lower center', ncol=3, bbox_to_anchor=(0.53, -0.04), title='Marker area',
               title_fontsize=7, fontsize=7, columnspacing=1.6)
    fig.tight_layout(rect=(0.02, 0.12, 1, 1), h_pad=1.2)
    save(fig, 'fig7_confidence_reliability.png')



def fig_flowchart():
    """System architecture diagram (Fig. 1): components and data flow only."""
    from matplotlib.patches import FancyBboxPatch, FancyArrowPatch, Rectangle
    W, H = 6.0, 3.32
    INK, INK2 = '#0b0b0b', '#3f3e3b'
    NUM, VIS, FUS, LLM, DAT = '#2a78d6', '#eb6834', '#138a5f', '#4a3aa7', '#77766f'
    FILL = {NUM: '#ffffff', VIS: '#ffffff', FUS: '#e6f6ef', LLM: '#f3f1fb', DAT: '#f3f2ee'}
    LANE = {NUM: '#eef4fd', VIS: '#fdf2ec'}

    with plt.rc_context({'font.family': 'DejaVu Sans', 'mathtext.fontset': 'dejavusans'}):
        fig = plt.figure(figsize=(W, H))
        ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(0, W); ax.set_ylim(0, H); ax.axis('off')
        blocks = []

        def lane(y0, y1, color, text):
            ax.add_patch(Rectangle((0.98, y0), 3.0, y1 - y0, fc=LANE[color], ec='none', zorder=0))
            ax.text(1.05, y1 - 0.11, text, ha='left', va='center', fontsize=6.4, fontweight='bold',
                    color=color, zorder=1)

        def block(x0, y0, x1, y1, color, title, sub=None, dashed=False, tsize=7.0, ssize=6.0):
            ax.add_patch(FancyBboxPatch((x0, y0), x1 - x0, y1 - y0, boxstyle='round,pad=0,rounding_size=0.06',
                                        fc=FILL[color], ec=color, lw=1.2, ls=(0, (4, 2)) if dashed else '-', zorder=2))
            cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
            texts = []
            if sub:
                nt, ns = title.count('\n') + 1, sub.count('\n') + 1
                lh_t, lh_s, g = tsize / 72 * 1.18, ssize / 72 * 1.22, 0.06
                tot = nt * lh_t + g + ns * lh_s
                top = cy + tot / 2
                texts.append(ax.text(cx, top - nt * lh_t / 2, title, ha='center', va='center', fontsize=tsize,
                                     fontweight='bold', color=INK, zorder=3, linespacing=1.15))
                texts.append(ax.text(cx, top - nt * lh_t - g - ns * lh_s / 2, sub, ha='center', va='center',
                                     fontsize=ssize, color=INK2, zorder=3, linespacing=1.2))
            else:
                texts.append(ax.text(cx, cy, title, ha='center', va='center', fontsize=tsize, fontweight='bold',
                                     color=INK, zorder=3, linespacing=1.15))
            blocks.append(((x0, y0, x1, y1), texts, title))
            return x0, y0, x1, y1

        def arrow(p0, p1, color='#55544f', dashed=False):
            ax.add_patch(FancyArrowPatch(p0, p1, arrowstyle='-|>', mutation_scale=8.5, lw=1.0, color=color,
                                         ls=(0, (3, 2)) if dashed else '-', zorder=1, shrinkA=0, shrinkB=0))

        def tag(x, y, s, color=INK2):
            ax.text(x, y, s, ha='center', va='center', fontsize=6.0, color=color, zorder=4,
                    bbox=dict(boxstyle='round,pad=0.1', fc='white', ec='none', alpha=0.95))

        # lanes ------------------------------------------------------------------
        NY0, NY1 = 1.86, 3.22          # numerical lane
        VY0, VY1 = 0.42, 1.78          # visual lane
        lane(NY0, NY1, NUM, 'NUMERICAL BRANCH')
        lane(VY0, VY1, VIS, 'VISUAL BRANCH')
        ny, vy = (NY0 + NY1) / 2 - 0.06, (VY0 + VY1) / 2 - 0.06   # block centre lines

        # blocks -----------------------------------------------------------------
        D = block(0.06, 0.95, 0.86, 2.70, DAT, 'C-MAPSS\nsensor\ndata', 'FD001–FD004\n3 settings\n21 sensors')
        P = block(1.08, ny - 0.36, 1.92, ny + 0.36, NUM, 'Preprocessing', None, tsize=6.5)
        M = block(2.12, ny - 0.36, 3.86, ny + 0.36, NUM, 'Numerical RUL models',
                  'Linear regression · Random forest\nXGBoost · BiLSTM · TCN · MS-TCN\nMS-TCN + BiLSTM · Transformer', ssize=5.8)
        R = block(1.08, vy - 0.36, 1.92, vy + 0.36, VIS, 'Sensor-plot\nrenderer', None)
        V = block(2.12, vy - 0.36, 3.02, vy + 0.36, VIS, 'VLM\nobserver', 'Qwen2.5-VL-\n7B-Instruct')
        S = block(3.16, vy - 0.36, 3.86, vy + 0.36, VIS, 'Schema\ncheck', 'validate,\nre-prompt', ssize=5.7)
        F = block(4.12, 0.94, 5.02, 2.62, FUS, 'Deterministic\nconstrained\nfusion', 'bounded VLM\nreweighting')
        O = block(5.24, 1.42, 5.96, 2.22, FUS, 'RUL\nestimate', '+ confidence')
        L = block(4.12, 0.0, 5.02, 0.66, LLM, 'Evidence-\nanalysis LLM', 'Qwen2.5-7B-\nInstruct', dashed=True, tsize=6.6, ssize=5.8)
        C = block(5.24, 0.06, 5.96, 0.60, LLM, 'Accept /\ncaution /\nreject', None, dashed=True, tsize=6.1)

        # flow -------------------------------------------------------------------
        arrow((D[2], ny), (P[0], ny)); arrow((D[2], vy), (R[0], vy))
        arrow((P[2], ny), (M[0], ny))
        arrow((R[2], vy), (V[0], vy)); arrow((V[2], vy), (S[0], vy))
        arrow((M[2], ny), (F[0], ny)); arrow((S[2], vy), (F[0], vy))
        arrow((F[2], (O[1] + O[3]) / 2), (O[0], (O[1] + O[3]) / 2))
        arrow(((F[0] + F[2]) / 2, F[1]), ((L[0] + L[2]) / 2, L[3]), LLM, dashed=True)
        arrow((L[2], 0.33), (C[0], 0.33), LLM, dashed=True)
        ax.text((F[0] + F[2]) / 2 + 0.07, (F[1] + L[3]) / 2, 'commentary only', ha='left', va='center', fontsize=5.9, style='italic', color=LLM, zorder=4)

        # fit check --------------------------------------------------------------
        fig.canvas.draw(); r = fig.canvas.get_renderer(); inv = ax.transData.inverted()
        bad = []
        for (x0, y0, x1, y1), texts, title in blocks:
            for t in texts:
                bb = inv.transform(t.get_window_extent(r))
                if bb[0][0] < x0 + 0.03 or bb[1][0] > x1 - 0.03 or bb[0][1] < y0 + 0.02 or bb[1][1] > y1 - 0.02:
                    bad.append(title.replace('\n', ' '))
        print('text-fit problems:', bad or 'none')
        fig.savefig(os.path.join(OUT, 'fig0_system_flowchart.png'), dpi=300, bbox_inches='tight', pad_inches=0.03)
        fig.savefig(os.path.join(OUT, 'fig0_system_flowchart.pdf'), bbox_inches='tight', pad_inches=0.03)
        plt.close(fig)
        print('wrote fig0_system_flowchart.png / .pdf')


# ----------------------------------------------------------------------------- run
for make in (fig_rmse, fig_skill, fig_contrib, fig_helphurt, fig_quality,
             fig_regime, fig_radar, fig_agreement, fig_confidence, fig_flowchart):
    make()

zip_path = OUT + ".zip"
with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
    for f in sorted(os.listdir(OUT)):
        z.write(os.path.join(OUT, f), os.path.join("figures", f))
print(f"\nAll figures saved in {OUT} and zipped to {zip_path}")
try:
    from google.colab import files
    files.download(zip_path)
except ImportError:
    # Only reached outside Colab (e.g. plain local Python), where google.colab does not
    # exist. The figures are still saved in OUT; only the browser download is skipped.
    pass
